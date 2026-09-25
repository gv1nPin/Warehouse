"""Журнал операций (operation_history) — только администратор."""

from __future__ import annotations

import json
from datetime import datetime

from django.contrib import messages
from django.shortcuts import redirect, render
from django.views.decorators.http import require_GET

from warehouse.common import PermissionName
from warehouse.common.exceptions import BusinessError

from ..auth import can, employee_required
from ..services import history_service
from ._helpers import bll_call, bll_err, bll_ok, eid


def _parse_dt(raw: str | None) -> datetime | None:
    if not raw:
        return None
    raw = raw.strip()
    for fmt in ("%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _details_preview(details: dict | None, limit: int = 120) -> str:
    if not details:
        return "—"
    try:
        text = json.dumps(details, ensure_ascii=False, default=str)
    except TypeError:
        text = str(details)
    if len(text) > limit:
        return text[: limit - 1] + "…"
    return text


@require_GET
@employee_required
def operations_list_view(request):
    """Список записей аудита. Право: employee:manage."""
    if not can(request, PermissionName.EMPLOYEE_MANAGE):
        messages.error(request, "Недостаточно прав для просмотра журнала операций")
        return redirect("home")

    employee_id = eid(request)
    q_employee = request.GET.get("employee_id") or ""
    q_type = (request.GET.get("operation_type") or "").strip()
    q_entity = (request.GET.get("entity_name") or "").strip()
    q_entity_id = request.GET.get("entity_id") or ""
    q_since = request.GET.get("since") or ""
    q_until = request.GET.get("until") or ""
    try:
        limit = min(int(request.GET.get("limit") or 100), 500)
    except ValueError:
        limit = 100
    try:
        offset = max(int(request.GET.get("offset") or 0), 0)
    except ValueError:
        offset = 0

    actor_employee_id = int(q_employee) if q_employee.isdigit() else None
    entity_id = int(q_entity_id) if q_entity_id.isdigit() else None
    since = _parse_dt(q_since)
    until = _parse_dt(q_until)

    filters: dict = {"operation_types": [], "entity_names": []}
    rows = []
    try:
        bll_call("OperationHistoryService.list_filters", request)
        filters = history_service().list_filters(employee_id)
        bll_ok("OperationHistoryService.list_filters")

        bll_call(
            "OperationHistoryService.list_operations",
            request,
            operation_type=q_type or None,
            entity_name=q_entity or None,
            limit=limit,
            offset=offset,
        )
        items = history_service().list_operations(
            employee_id,
            actor_employee_id=actor_employee_id,
            operation_type=q_type or None,
            entity_name=q_entity or None,
            entity_id=entity_id,
            since=since,
            until=until,
            limit=limit,
            offset=offset,
        )
        bll_ok("OperationHistoryService.list_operations", count=len(items))
        for op in items:
            created = op.created_at
            if hasattr(created, "strftime"):
                created_s = created.strftime("%d.%m.%Y %H:%M:%S")
            else:
                created_s = str(created)
            rows.append(
                {
                    "id": op.id,
                    "created_at": created_s,
                    "employee_id": op.employee_id,
                    "employee_name": op.employee_name,
                    "operation_type": op.operation_type,
                    "entity_name": op.entity_name,
                    "entity_id": op.entity_id,
                    "details_preview": _details_preview(op.details),
                    "details_json": json.dumps(op.details, ensure_ascii=False, default=str)
                    if op.details
                    else "",
                }
            )
    except BusinessError as exc:
        bll_err("OperationHistoryService.list_operations", request, exc)
        messages.error(request, getattr(exc, "message", str(exc)))

    return render(
        request,
        "web/pages/operations.html",
        {
            "rows": rows,
            "operation_types": filters.get("operation_types") or [],
            "entity_names": filters.get("entity_names") or [],
            "q_employee_id": q_employee,
            "q_operation_type": q_type,
            "q_entity_name": q_entity,
            "q_entity_id": q_entity_id,
            "q_since": q_since,
            "q_until": q_until,
            "limit": limit,
            "offset": offset,
            "next_offset": offset + limit if len(rows) >= limit else None,
            "prev_offset": max(offset - limit, 0) if offset > 0 else None,
        },
    )


@require_GET
@employee_required
def operation_detail_view(request, operation_id: int):
    """Карточка одной записи журнала."""
    if not can(request, PermissionName.EMPLOYEE_MANAGE):
        messages.error(request, "Недостаточно прав")
        return redirect("home")

    employee_id = eid(request)
    try:
        bll_call("OperationHistoryService.get_operation", request, operation_id=operation_id)
        op = history_service().get_operation(employee_id, operation_id)
        bll_ok("OperationHistoryService.get_operation", op)
    except BusinessError as exc:
        bll_err("OperationHistoryService.get_operation", request, exc)
        messages.error(request, getattr(exc, "message", str(exc)))
        return redirect("operations")

    created = op.created_at
    created_s = created.strftime("%d.%m.%Y %H:%M:%S") if hasattr(created, "strftime") else str(created)
    details_pretty = (
        json.dumps(op.details, ensure_ascii=False, indent=2, default=str) if op.details else "—"
    )
    return render(
        request,
        "web/pages/operation_detail.html",
        {
            "op": {
                "id": op.id,
                "created_at": created_s,
                "employee_id": op.employee_id,
                "employee_name": op.employee_name,
                "operation_type": op.operation_type,
                "entity_name": op.entity_name,
                "entity_id": op.entity_id,
                "details_pretty": details_pretty,
            }
        },
    )


@require_GET
@employee_required
def operations_export_view(request):
    """Выгрузка журнала в Excel (.xlsx). Те же фильтры, что у списка."""
    if not can(request, PermissionName.EMPLOYEE_MANAGE):
        messages.error(request, "Недостаточно прав для экспорта журнала")
        return redirect("home")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        messages.error(
            request,
            "Для экспорта в Excel установите openpyxl: pip install openpyxl",
        )
        return redirect("operations")

    employee_id = eid(request)
    q_employee = request.GET.get("employee_id") or ""
    q_type = (request.GET.get("operation_type") or "").strip()
    q_entity = (request.GET.get("entity_name") or "").strip()
    q_entity_id = request.GET.get("entity_id") or ""
    q_since = request.GET.get("since") or ""
    q_until = request.GET.get("until") or ""

    actor_employee_id = int(q_employee) if q_employee.isdigit() else None
    entity_id = int(q_entity_id) if q_entity_id.isdigit() else None
    since = _parse_dt(q_since)
    until = _parse_dt(q_until)

    # Экспорт: до 5000 строк (верхний предел репозитория — 500 за раз → пагинация)
    all_rows = []
    offset = 0
    page = 500
    try:
        while True:
            bll_call(
                "OperationHistoryService.list_operations",
                request,
                export=True,
                limit=page,
                offset=offset,
            )
            batch = history_service().list_operations(
                employee_id,
                actor_employee_id=actor_employee_id,
                operation_type=q_type or None,
                entity_name=q_entity or None,
                entity_id=entity_id,
                since=since,
                until=until,
                limit=page,
                offset=offset,
            )
            bll_ok("OperationHistoryService.list_operations", count=len(batch), export=True)
            all_rows.extend(batch)
            if len(batch) < page or len(all_rows) >= 5000:
                break
            offset += page
    except BusinessError as exc:
        bll_err("OperationHistoryService.list_operations", request, exc)
        messages.error(request, getattr(exc, "message", str(exc)))
        return redirect("operations")

    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал операций"

    headers = [
        "ID",
        "Время",
        "ID сотрудника",
        "Сотрудник",
        "Тип операции",
        "Сущность",
        "ID сущности",
        "Детали (JSON)",
    ]
    ws.append(headers)
    header_font = Font(bold=True)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for op in all_rows:
        created = op.created_at
        if hasattr(created, "strftime"):
            created_s = created.strftime("%d.%m.%Y %H:%M:%S")
        else:
            created_s = str(created)
        details = ""
        if op.details:
            try:
                details = json.dumps(op.details, ensure_ascii=False, default=str)
            except TypeError:
                details = str(op.details)
        ws.append(
            [
                op.id,
                created_s,
                op.employee_id,
                op.employee_name,
                op.operation_type,
                op.entity_name,
                op.entity_id if op.entity_id is not None else "",
                details,
            ]
        )

    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from django.http import HttpResponse

    widths = [8, 20, 14, 28, 22, 18, 12, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"operation_history_{stamp}.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
